#!/usr/bin/env python3
"""Prepare Prism XML data and execute GraphPad's native Windows scripts.

Standard library only, except optional openpyxl for XLSX input.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import os
from pathlib import Path
import subprocess
import sys
import time
import uuid
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime
from decimal import Decimal, InvalidOperation

NS = "http://graphpad.com/prism/Prism.htm"
ET.register_namespace("", NS)
ET.register_namespace("xsl", "http://www.w3.org/1999/XSL/Transform")
PRESETS = {
    "column-scatter": ("column", "Column scatter.pzt"),
    "box-whisker": ("column", "Box and whiskers graph.pzt"),
    "scatter-bars": ("column", "Scatter plot with bars.pzt"),
    "xy-line": ("xy", "Spaghetti plot.pzt"),
}


def tag(name: str) -> str:
    return f"{{{NS}}}{name}"


def read_text(path: Path) -> str:
    raw = path.read_bytes()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            pass
    raise ValueError(f"不能解码输入文件: {path}")


def number(value, row: int, column: int):
    if value is None or isinstance(value, str) and not value.strip():
        return None
    if isinstance(value, bool):
        raise ValueError(f"第 {row} 行第 {column} 列是布尔值，需明确数值含义")
    try:
        result = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"第 {row} 行第 {column} 列不是数值: {value!r}") from None
    if not math.isfinite(result):
        raise ValueError(f"第 {row} 行第 {column} 列含 NaN/Inf，缺测请用空值")
    # Keep integers and numeric strings exact instead of rounding through float.
    return str(value).strip()


def load_data(path: Path, sheet: str | None):
    suffix = path.suffix.lower()
    if suffix == ".json":
        payload = json.loads(read_text(path))
        if not isinstance(payload, dict) or not isinstance(payload.get("columns"), list) or not isinstance(payload.get("rows"), list):
            raise ValueError('JSON 需要 {"columns": [...], "rows": [[...]]}')
        matrix = [payload["columns"], *payload["rows"]]
    elif suffix in (".csv", ".tsv", ".txt"):
        contents = read_text(path)
        delimiter = "\t" if suffix == ".tsv" else ","
        if suffix == ".txt":
            delimiter = csv.Sniffer().sniff(contents[:8192], delimiters=",\t;").delimiter
        matrix = list(csv.reader(io.StringIO(contents), delimiter=delimiter))
    elif suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ValueError("XLSX 需要 openpyxl；请用 uv run --with openpyxl python 运行") from None
        workbook = load_workbook(path, data_only=False, read_only=True)
        values = load_workbook(path, data_only=True, read_only=True)
        try:
            if sheet:
                if sheet not in workbook.sheetnames:
                    raise ValueError(f"没有工作表 {sheet!r}；可用: {workbook.sheetnames}")
                selected = sheet
            else:
                nonempty = [ws.title for ws in workbook if any(any(c.value is not None for c in row) for row in ws.iter_rows())]
                if len(nonempty) != 1:
                    raise ValueError(f"需用 --sheet 选择工作表；非空表: {nonempty}")
                selected = nonempty[0]
            matrix = []
            for source_row, cached_row in zip(workbook[selected].iter_rows(), values[selected].iter_rows()):
                row = []
                for source, cached in zip(source_row, cached_row):
                    if source.data_type == "f" and cached.value is None:
                        raise ValueError(f"公式 {selected}!{source.coordinate} 没有缓存结果，请用 Excel 重算并保存")
                    row.append(cached.value)
                matrix.append(row)
            # Ignore unused trailing worksheet columns, never an interior blank.
            used = max((i + 1 for row in matrix for i, value in enumerate(row) if value is not None), default=0)
            matrix = [row[:used] for row in matrix]
        finally:
            workbook.close()
            values.close()
    else:
        raise ValueError("输入支持 .csv、.tsv、.txt、.xlsx、.json")
    while matrix and all(v is None or isinstance(v, str) and not v.strip() for v in matrix[-1]):
        matrix.pop()
    if len(matrix) < 2 or not isinstance(matrix[0], (list, tuple)):
        raise ValueError("输入需要表头和至少一行数据")
    headers = [str(value).strip() if value is not None else "" for value in matrix[0]]
    if not all(headers) or len(set(headers)) != len(headers):
        raise ValueError("表头必须非空且不能重复")
    rows = []
    for index, row in enumerate(matrix[1:], 2):
        if not isinstance(row, (list, tuple)) or len(row) != len(headers):
            raise ValueError(f"第 {index} 行长度与表头不一致")
        rows.append([number(value, index, column + 1) for column, value in enumerate(row)])
    return headers, rows


def prism_candidates():
    roots = [Path(os.environ.get(key, fallback)) / "GraphPad" for key, fallback in (
        ("ProgramFiles", "C:/Program Files"), ("ProgramFiles(x86)", "C:/Program Files (x86)"))]
    found = []
    for root in roots:
        if root.is_dir():
            found.extend(root.glob("*/prism.exe"))
            found.extend(root.glob("*/*/prism_CN.exe"))
    return list(dict.fromkeys(p.resolve() for p in found))


def find_prism(explicit: str | None) -> Path:
    preferred = explicit or os.environ.get("PRISM_EXE")
    if preferred:
        path = Path(preferred).resolve()
        if not path.is_file():
            raise ValueError(f"Prism 可执行文件不存在: {path}")
        return path
    found = prism_candidates()
    if not found:
        raise ValueError("未找到 Prism，请指定 --prism-exe 或 PRISM_EXE")
    return found[0]


def parse_template(path: Path):
    tree = ET.parse(path)
    root = tree.getroot()
    project = root if root.tag == tag("GraphPadPrismFile") else root.find(f".//{tag('GraphPadPrismFile')}")
    if project is None:
        raise ValueError("不是受支持的 Prism XML 项目")
    by_id = {table.get("ID"): table for table in project.findall(tag("Table"))}
    sequence = project.find(tag("TableSequence"))
    tables = [by_id[ref.get("ID")] for ref in sequence if ref.get("ID") in by_id] if sequence is not None else list(by_id.values())
    if not tables:
        raise ValueError("模板没有可编辑数据表（HugeTable/Table1024 需另行转换）")
    return tree, tables


def inspect_template(path: Path):
    _, tables = parse_template(path)
    return [{
        "index": i, "id": table.get("ID"), "title": table.findtext(tag("Title")),
        "type": table.get("TableType"), "x_format": table.get("XFormat"),
        "y_format": table.get("YFormat"),
        "datasets": len(table.findall(tag("YColumn"))),
        "subcolumns": [len(column.findall(tag("Subcolumn"))) for column in table.findall(tag("YColumn"))],
    } for i, table in enumerate(tables, 1)]


def replace_values(column, values, title):
    title_node = column.find(tag("Title"))
    if title_node is None:
        title_node = ET.Element(tag("Title"))
        column.insert(0, title_node)
    title_node.clear()
    title_node.text = title
    subcolumns = column.findall(tag("Subcolumn"))
    if len(subcolumns) != 1:
        raise ValueError("模板列必须且只能有一个 Subcolumn")
    subcolumns[0].clear()
    for value in values:
        ET.SubElement(subcolumns[0], tag("d")).text = value


def prepare(template: Path, destination: Path, headers, rows, layout, table_index):
    tree, tables = parse_template(template)
    if table_index is None:
        if len(tables) != 1:
            raise ValueError("模板包含多张数据表，需明确指定 --table-index，并核对 --graph-index 的数据引用")
        table_index = 1
    if not 1 <= table_index <= len(tables):
        raise ValueError("--table-index 超出模板范围")
    table = tables[table_index - 1]
    expected = "OneWay" if layout == "column" else "XY"
    if table.get("TableType") != expected:
        raise ValueError(f"需要 {expected} 模板，实际为 {table.get('TableType')}")
    if table.get("YFormat", "replicates") not in ("replicates", ""):
        raise ValueError(f"不支持模板的 YFormat={table.get('YFormat')}，需要原始单列观测值")
    columns = table.findall(tag("YColumn"))
    offset = 1 if layout == "xy" else 0
    dataset_count = len(headers) - offset
    if dataset_count < 1 or dataset_count > len(columns):
        raise ValueError(f"输入有 {dataset_count} 个数据集，模板容量为 {len(columns)}；请提供匹配模板")
    if any(not any(row[i] is not None for row in rows) for i in range(offset, len(headers))):
        raise ValueError("输入有完全空的数据集，需先明确其绘图含义")
    if any(len(column.findall(tag("Subcolumn"))) != 1 for column in columns):
        raise ValueError("模板含重复或误差子列，不支持当前单 Y 列输入")
    if layout == "xy":
        if table.get("XFormat") not in ("numbers", "numeric"):
            raise ValueError(f"需要数值 X 模板，实际 XFormat={table.get('XFormat')}")
        if any(row[0] is None for row in rows):
            raise ValueError("XY 输入的 X 列不能空")
        xcolumn = table.find(tag("XColumn"))
        if xcolumn is None:
            raise ValueError("XY 模板缺少 XColumn")
        replace_values(xcolumn, [row[0] for row in rows], headers[0])
    for i, column in enumerate(columns):
        values = [row[i + offset] for row in rows] if i < dataset_count else []
        replace_values(column, values, headers[i + offset] if i < dataset_count else "")
    # Demo descriptions and sticky notes are data-table annotations, not graphs.
    for note in list(table.findall(tag("FloatingNote"))):
        table.remove(note)
    table_title = table.find(tag("Title"))
    if table_title is not None:
        table_title.clear()
        table_title.text = "Input data"
    tree.write(destination, encoding="utf-8", xml_declaration=True)
    return {"table_index": table_index, "dataset_count": dataset_count, "row_count": len(rows),
            "unused_template_datasets": len(columns) - dataset_count}


def quoted(value) -> str:
    value = str(value)
    if any(character in value for character in ('"', "\n", "\r", "\0")):
        raise ValueError("Prism 参数不能包含双引号、换行或 NUL")
    return '"' + value.replace("%", "%%") + '"'


def write_script(path: Path, lines):
    # Prism treats a UTF-8 BOM as part of the first command name.
    path.write_bytes(("\r\n".join(lines) + "\r\n").encode("utf-8"))


def completion_lines(marker: Path):
    return [f"OpenOutput {quoted(marker)}", 'WText "done"', "CloseOutput"]


def execute(prism: Path, script: Path, marker: Path, timeout: float):
    if os.name != "nt":
        raise ValueError("此处理器的自动执行入口仅支持 Windows")
    startup = subprocess.STARTUPINFO()
    startup.dwFlags |= subprocess.STARTF_USESHOWWINDOW
    startup.wShowWindow = 0
    started = time.monotonic()
    process = subprocess.Popen([str(prism), "@" + str(script)], startupinfo=startup)
    exit_observed = None
    while time.monotonic() - started < timeout:
        if marker.is_file() and marker.read_text(encoding="utf-8-sig", errors="replace").strip().strip('"') == "done":
            return {"pid": process.pid, "seconds": round(time.monotonic() - started, 2)}
        exit_code = process.poll()
        if exit_code is not None:
            exit_observed = exit_observed or time.monotonic()
            # Allow forwarding to an existing Prism instance before diagnosing.
            if time.monotonic() - exit_observed > min(10, timeout / 2):
                raise RuntimeError(f"Prism 退出码 {exit_code}，但未生成完成标记；检查 {script.with_suffix('.log')}，或在 Prism 内手动运行 {script}")
        time.sleep(0.25)
    raise RuntimeError(f"Prism 执行超过 {timeout:g} 秒；保留进程 {process.pid} 和运行目录供诊断，不自动终止用户窗口")


def validate_export(path: Path, fmt: str):
    if not path.is_file() or path.stat().st_size == 0:
        raise RuntimeError(f"Prism 没有生成导出文件: {path}")
    with path.open("rb") as stream:
        prefix = stream.read(4096)
    signatures = {
        "pdf": lambda b: b.startswith(b"%PDF-"),
        "png": lambda b: b.startswith(b"\x89PNG\r\n\x1a\n"),
        "tif": lambda b: b.startswith((b"II*\x00", b"MM\x00*")),
        "eps": lambda b: b.startswith(b"%!PS"),
        "emf": lambda b: len(b) >= 44 and b[40:44] == b" EMF",
        "svg": lambda b: b"<svg" in b,
    }
    if not signatures[fmt](prefix):
        raise RuntimeError(f"导出文件的格式签名与 {fmt} 不一致: {path}")


def verify_project(path: Path, headers, rows, layout, table_index):
    _, tables = parse_template(path)
    table = tables[table_index - 1]
    columns = table.findall(tag("YColumn"))[:len(headers) - (layout == "xy")]
    if layout == "xy":
        columns.insert(0, table.find(tag("XColumn")))
    for index, column in enumerate(columns):
        if column is None or column.findtext(tag("Title"), "") != headers[index]:
            raise RuntimeError(f"输出项目的第 {index + 1} 列标题不符")
        values = [cell.text for cell in column.find(tag("Subcolumn"))]
        expected = [row[index] for row in rows]
        # Prism can omit trailing empty cells and normalize numeric formatting.
        while expected and expected[-1] is None:
            expected.pop()
        while values and values[-1] is None:
            values.pop()
        try:
            mismatch = len(values) != len(expected) or any(
                (actual is None) != (original is None) or
                actual is not None and Decimal(actual) != Decimal(original)
                for actual, original in zip(values, expected)
            )
        except InvalidOperation:
            mismatch = True
        if mismatch:
            raise RuntimeError(f"输出项目的第 {index + 1} 列数据与输入不符")


def validate_prism_project(path: Path):
    """Check that Prism saved a real modern project, rather than renamed XML."""
    if not path.is_file() or not zipfile.is_zipfile(path):
        raise RuntimeError(f"不是有效的 .prism 原生项目: {path}")
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise RuntimeError(f".prism 项目压缩包校验失败: {path}")
        try:
            document = json.loads(archive.read("document.json"))
            sheets = document["sheets"]
            for uid in sheets["data"]:
                archive.read(f"data/sheets/{uid}/sheet.json")
            for uid in sheets["graphs"]:
                archive.read(f"graphs/{uid}/sheet.json")
                if not archive.read(f"graphs/{uid}/data.bin"):
                    raise ValueError("图形配置为空")
            for uid in sheets["analyses"]:
                archive.read(f"analyses/{uid}/parameters.json")
                if not json.loads(archive.read(f"analyses/{uid}/results.json")):
                    raise ValueError("原生分析结果为空")
        except (KeyError, ValueError) as error:
            raise RuntimeError(f".prism 项目缺少完整数据、图形或分析内容: {path}") from error


def save_manifest(path, payload):
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def plot(args):
    source = Path(args.input).resolve()
    headers, rows = load_data(source, args.sheet)
    prism = None if args.prepare_only else find_prism(args.prism_exe)
    if args.template:
        template = Path(args.template).resolve()
        if not template.is_file() or template.suffix.lower() not in (".pzfx", ".pzt", ".pzf", ".prism"):
            raise ValueError("模板需要存在的 .pzfx、.pzt、.pzf 或 .prism 文件")
    else:
        if args.prepare_only:
            raise ValueError("--prepare-only 需要 --template 指向 .pzfx")
        preset = args.preset or ("column-scatter" if args.layout == "column" else "xy-line")
        layout, filename = PRESETS[preset]
        if layout != args.layout:
            raise ValueError(f"预设 {preset} 需要 --layout {layout}")
        template = prism.parent / "Portfolio" / "Graphs to explore" / filename
        if not template.is_file():
            raise ValueError(f"本机没有预设模板 {template}，请指定 --template")
    if args.prepare_only and template.suffix.lower() != ".pzfx":
        raise ValueError("非 .pzfx 模板需由 Prism 转换，--prepare-only 仅接受 .pzfx")
    output = Path(args.output).resolve() if args.output else source.parent / "prism-output"
    run = output / (source.stem + "-" + datetime.now().strftime("%Y%m%d-%H%M%S-") + uuid.uuid4().hex[:8])
    run.mkdir(parents=True, exist_ok=False)
    work = run / "_work"
    work.mkdir()
    manifest_path = run / "manifest.json"
    manifest = {"status": "preparing", "source": str(source), "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                "template": str(template), "prism_exe": str(prism) if prism else None,
                "layout": args.layout, "headers": headers, "graph_index": args.graph_index,
                "project_format": args.project_format, "color_scheme": args.color_scheme,
                "run_directory": str(run), "work_directory": str(work),
                "warnings": ["继承模板的图形样式、统计和坐标设置；交付前必须检查预览"]}
    save_manifest(manifest_path, manifest)
    try:
        with (work / "normalized-data.csv").open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(headers)
            writer.writerows(rows)
        if template.suffix.lower() != ".pzfx":
            converted = work / "template.pzfx"
            marker = work / "template-finished.txt"
            script = work / "convert-template.pzc"
            write_script(script, ["CreateLog", f"Open {quoted(template)}", f"Save {quoted(converted)}", "Close", *completion_lines(marker)])
            execute(prism, script, marker, args.timeout)
            template = converted
        prepared = work / "prepared.pzfx"
        manifest["data"] = prepare(template, prepared, headers, rows, args.layout, args.table_index)
        project = run / f"{source.stem}.{args.project_format}"
        verification_project = work / "verification.pzfx" if args.project_format == "prism" else project
        script = work / "plot.pzc"
        marker = work / "finished.txt"
        lines = ["CreateLog", f"Open {quoted(prepared)}", "RecalcAll"]
        results_export = None
        if args.results_index is not None:
            results_export = run / "statistics.txt"
            lines += [f"GoTo R {args.results_index}", f"ExportTable {quoted(results_export)}"]
        lines.append(f"GoTo G {args.graph_index}")
        if args.color_scheme is not None:
            lines.append(f"ApplyColorScheme {quoted(args.color_scheme)}")
        for command, value in (("SetGraphTitle", args.title), ("SetAxisTitle X", args.x_title), ("SetAxisTitle Y", args.y_title)):
            if value is not None:
                lines.append(f"{command} {quoted(value)}")
        for axis in ("x", "y"):
            for bound, keyword in (("min", "bottom"), ("max", "top")):
                value = getattr(args, f"{axis}_{bound}")
                if value is not None:
                    lines.append(f"SetAxisLimits {axis.upper()} {keyword} {value:g}")
            interval = getattr(args, f"{axis}_interval")
            if interval is not None:
                lines.append(f"SetAxisLimits {axis.upper()} Interval {interval:g}")
        exports = []
        commands = {"pdf": "ExportPDF", "png": "ExportPNG", "tif": "ExportTIF", "eps": "ExportEPS", "emf": "ExportEMF", "svg": "ExportSVG"}
        for fmt in dict.fromkeys(args.formats):
            path = run / f"{source.stem}.{fmt}"
            options = f" {args.dpi} RGBhi LZW" if fmt == "tif" else f" {args.dpi}" if fmt == "png" else " RGB" if fmt in ("pdf", "eps") else ""
            lines.append(f"{commands[fmt]} {quoted(path)}{options}")
            exports.append((path, fmt))
        # Keep XML for exact input verification; the primary native project is
        # saved by Prism itself, preserving its analysis results and graph state.
        lines.append(f"Save {quoted(verification_project)}")
        if project != verification_project:
            lines.append(f"Save {quoted(project)}")
        lines += ["Close", *completion_lines(marker)]
        write_script(script, lines)
        manifest.update({"script": str(script), "prepared_project": str(prepared), "project": str(project),
                         "verification_project": str(verification_project), "exports": [str(path) for path, _ in exports]})
        if results_export is not None:
            manifest.update({"results_index": args.results_index, "results_export": str(results_export)})
        if args.prepare_only:
            verify_project(prepared, headers, rows, args.layout, manifest["data"]["table_index"])
            manifest["status"] = "prepared_only"
        else:
            manifest["execution"] = execute(prism, script, marker, args.timeout)
            for path, fmt in exports:
                validate_export(path, fmt)
            if results_export is not None and (not results_export.is_file() or results_export.stat().st_size == 0):
                raise RuntimeError(f"Prism 没有生成非空结果表: {results_export}")
            verify_project(verification_project, headers, rows, args.layout, manifest["data"]["table_index"])
            if args.project_format == "prism":
                validate_prism_project(project)
            manifest.update({"deliverable_project": str(project),
                             "deliverable_project_sha256": hashlib.sha256(project.read_bytes()).hexdigest()})
            manifest["status"] = "exported_preview_pending"
        save_manifest(manifest_path, manifest)
        print(json.dumps(manifest, ensure_ascii=False, indent=2))
    except Exception as error:
        manifest.update({"status": "failed", "error": str(error)})
        save_manifest(manifest_path, manifest)
        raise RuntimeError(f"{error}\n诊断记录: {manifest_path}") from error


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    doctor = sub.add_parser("doctor", help="只检查 Prism 安装，不启动程序")
    doctor.add_argument("--prism-exe")
    inspect = sub.add_parser("inspect-template")
    inspect.add_argument("--template", required=True)
    plotting = sub.add_parser("plot")
    plotting.add_argument("--input", required=True)
    plotting.add_argument("--sheet")
    plotting.add_argument("--layout", required=True, choices=("column", "xy"))
    plotting.add_argument("--template")
    plotting.add_argument("--preset", choices=tuple(PRESETS))
    plotting.add_argument("--prism-exe")
    plotting.add_argument("--output")
    plotting.add_argument("--table-index", type=int)
    plotting.add_argument("--graph-index", type=int, default=1)
    plotting.add_argument("--results-index", type=int, help="导出指定 Prism 原生 Results 表为 statistics.txt；序号从 1 开始")
    plotting.add_argument("--project-format", choices=("prism", "pzfx"), default="prism",
                          help="交付源项目格式；prism 需要 Prism 10+，旧版可选 pzfx")
    plotting.add_argument("--color-scheme", help="Prism 已有配色方案名称，例如 Colorblind Safe、Floral 或 Waves")
    plotting.add_argument("--title")
    plotting.add_argument("--x-title")
    plotting.add_argument("--y-title")
    plotting.add_argument("--formats", nargs="+", choices=("pdf", "png", "tif", "eps", "emf", "svg"), default=["pdf", "png"])
    plotting.add_argument("--dpi", type=int, choices=(100, 300, 600, 1200), default=300)
    plotting.add_argument("--timeout", type=float, default=120)
    plotting.add_argument("--prepare-only", action="store_true")
    for axis in ("x", "y"):
        for bound in ("min", "max"):
            plotting.add_argument(f"--{axis}-{bound}", type=float)
        plotting.add_argument(f"--{axis}-interval", type=float, help="主刻度间距，需大于零")
    args = parser.parse_args()
    try:
        if args.command == "doctor":
            prism = find_prism(args.prism_exe)
            print(json.dumps({"prism_exe": str(prism), "candidates": [str(p) for p in prism_candidates()],
                              "presets": {name: str(prism.parent / "Portfolio" / "Graphs to explore" / filename)
                                          for name, (_, filename) in PRESETS.items()},
                              "note": "路径检查完成；尚未验证 Prism 执行和导出"}, ensure_ascii=False, indent=2))
        elif args.command == "inspect-template":
            print(json.dumps(inspect_template(Path(args.template)), ensure_ascii=False, indent=2))
        else:
            if args.timeout <= 0 or args.graph_index < 1:
                raise ValueError("timeout 和 graph-index 必须大于零")
            if args.results_index is not None and args.results_index < 1:
                raise ValueError("results-index 必须大于零")
            for axis in ("x", "y"):
                low, high = getattr(args, f"{axis}_min"), getattr(args, f"{axis}_max")
                if any(value is not None and not math.isfinite(value) for value in (low, high)) or low is not None and high is not None and low >= high:
                    raise ValueError(f"{axis.upper()} 坐标范围必须有限且最小值小于最大值")
                interval = getattr(args, f"{axis}_interval")
                if interval is not None and (not math.isfinite(interval) or interval <= 0):
                    raise ValueError(f"{axis.upper()} 主刻度间距必须是大于零的有限数值")
            plot(args)
        return 0
    except (ValueError, RuntimeError, OSError, ET.ParseError, json.JSONDecodeError) as error:
        print(str(error), file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
